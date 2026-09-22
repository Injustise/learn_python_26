from docx import Document
from docx.shared import Cm, Pt
from docx.document import Document as Doc

# 创建 Word 文档对象
document = Document()
# 添加大标题
document.add_heading('Learn Python With Joy!', 0) 
# 添加段落，并创建该段落对象
para = document.add_paragraph('Python 是一门非常流行的编程语言，它')
# 添加文本块，并创建该文本块对象
run = para.add_run('简单')
#设置该文本块格式
run.bold = True
run.font.size = Pt(18)

para.add_run('而且')

run = para.add_run('优雅')
run.bold = True
run.font.size = Pt(18)
run.underline = True

para.add_run('。')

# 添加一级标题
document.add_heading('Heading level 1', level = 1)
# 添加带样式的段落
para = document.add_paragraph('Intense Quote', style = 'Intense Quote')
para.add_run('：Learning Knows No Bounds!')
# 添加无序列表（Bullet）
document.add_paragraph(
    'items in unordered list', style = 'List Bullet'
)
# 添加有序列表（Number）
document.add_paragraph(
    'first item in ordered list', style = 'List Number'
)
document.add_paragraph(
    'second item in ordered list', style = 'List Number'
)
# 添加图片
document.add_picture('D://code_py//learn_9-1//py文件操作//04_Word//明日香.jpg', width = Cm(5.2)) # 图片宽度固定为 5.2 厘米，高度按比例自动缩放

# 添加分节符
document.add_section() # 页面布局可以改变

# 数据容器完全由自己的喜好和数据结构需求决定
"""
不可变用元组
可变用列表
带键名用字典
"""
records = (
    ('张竞航', '男', '2008-04-30'),
    ('七七', '女', '2010-07-07')
)
# 添加表格，并创建该表格对象
table = document.add_table(rows = 1, cols = 3) # 创建一个 1 行 3 列的空表格作表头，后用 add_row() 动态追加，省去手动计算总行数
table.style = 'Light Grid'
hdr_cells = table.rows[0].cells # table.rows[0] 获取第一行的单元格集合
# word 的 cell.text 类比 excel 的 cell.value， 但 .text 一定是纯字符串
hdr_cells[0].text = '姓名'
hdr_cells[1].text = '性别'
hdr_cells[2].text = '出生日期'

for name, sex, bir, in records:
    row_cells = table.add_row().cells # table.add_row()：每次循环都在表格末尾新增一行空白行，并获取该行的单元格集合
    row_cells[0].text = name
    row_cells[1].text = sex
    row_cells[2].text = bir

# 添加分页符
document.add_page_break() # 页面布局不可以改变

# 保存文档
document.save('D://code_py//learn_9-1//py文件操作//04_Word//Learn Python.docx')


# 模板文件
students = [
    {
        'name' : '张竞航',
        'id' : '20262563',
        'date' : '2026年9月20日',
        'reason' : '感冒发烧',
        'duration' : '3天',
        'instructor' : '七七'
    }
]

for stu in students:
    # 获取文档对象
    doc = Document('D://code_py//learn_9-1//py文件操作//04_Word//请假条.docx')
    for para in doc.paragraphs: # 遍历段落
        if '{' not in para.text:
            continue
        """ 
        Word 都会把类似 {name} 这样的占位符拆成至少 2~3 个独立的 run，导致无法通过遍历 run 来锁定占位符
        for run in p.runs:
            if '{' not in run.text:
                continue
        """
        for key, value in stu.items():
            place_holder = f'{{{key}}}' # 内层 {key} 表示 key 的值， 外层 {{ 表示字符 {（转义） 
            if place_holder in para.text:
                para.text = para.text.replace(place_holder, value) # replace 替换占位符
                """
                当你执行 para.text = ...，其内部会自动执行两步操作：
                1.清空这个段落里所有的旧 run（相当于 para.clear()）。
                2.新建一个 run，把新的内容填进去（相当于 para.add_run(...)）。
                """
                # 故旧 run 和 新建的 run 格式出现不同
    doc.save(f'D://code_py//learn_9-1//py文件操作//04_Word//{stu['name']}请假条.docx')